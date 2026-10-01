#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
kci_kwater_review — K-water 소속 논문 분류 검증용 표 만들기 (사람이 눈으로 확인)

KCI 전체에서 affiliation=검색어(한국수자원공사·K-water 등)로 다시 조회해,
저자 소속 문자열 재확인(kci_affiliation_pattern)에서
  · 채택된 논문(accepted)  → 오탐(K-water가 아닌데 들어옴) 확인용
  · 탈락한 논문(rejected)  → 누락(K-water인데 빠짐) 확인용
을 모두 표로 남긴다. 소속 문자열별 빈도표도 만든다.

출력 (비공개 저장소에만 — KCI 이용 준수: 원천 데이터 재배포 금지):
  {KCI_RAW_DIR}/review/kwater_papers.csv        논문 단위 (판정·연도·학술지·제목·K-water 저자 소속·URL)
  {KCI_RAW_DIR}/review/kwater_affiliations.csv  소속 문자열 단위 (빈도·판정·걸린 검색어)
  {KCI_RAW_DIR}/review/kwater_review.xlsx       위 두 표 + 검증(O/X)·메모 칸 + 요약 (내 PC에서 검증용)
  {KCI_RAW_DIR}/review/README.md                보는 법
로그에는 건수만 출력 (제목·이름 미출력).
"""
import csv
import os
import sys
import time
from collections import Counter, defaultdict
from pathlib import Path

from kci_collect import KEY, KW_PATTERN, KW_QUERIES, fetch_year
from water_journals import load_registry

OUT = Path(os.environ.get("KCI_RAW_DIR", "private")) / "review"


def write_xlsx(path: Path, paper_rows: list, aff_rows: list, meta: dict):
    """검증용 엑셀: 안내·요약(수식)·소속 문자열·논문(탈락/채택). 노란 칸 = 사용자가 채우는 칸."""
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    F = Font(name="Arial", size=10)
    FB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    HEAD = PatternFill("solid", fgColor="17427C")
    FILL = PatternFill("solid", fgColor="FFF2B2")                # 사용자 입력 칸
    wb = Workbook()

    def table(ws, header, rows, widths, input_cols):
        ws.append(header)
        for r in rows:
            ws.append(list(r) + [None] * len(input_cols))          # 진짜 빈칸 (COUNTA가 세지 않게)
        for c in ws[1]:
            c.font, c.fill = FB, HEAD
            c.alignment = Alignment(vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = F
                c.alignment = Alignment(vertical="top", wrap_text=True)
        n = len(header)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        last = ws.max_row
        for j in range(n - len(input_cols) + 1, n + 1):
            col = get_column_letter(j)
            for i in range(2, last + 1):
                ws[f"{col}{i}"].fill = FILL
        vcol = get_column_letter(n - len(input_cols) + 1)       # 검증 칸: O/X 드롭다운
        dv = DataValidation(type="list", formula1='"O,X"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{vcol}2:{vcol}{max(last, 2)}")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(n)}{max(last, 1)}"

    # ── 안내 ──
    ws = wb.active
    ws.title = "안내"
    lines = [
        ("K-water 소속 논문 분류 검증표", ""),
        ("출처", "KCI(한국학술지인용색인) 데이터 활용 — 비공개 보관, 재배포 금지"),
        ("갱신", meta["updated"]),
        ("검색어 (KCI affiliation=)", ", ".join(KW_QUERIES)),
        ("재확인 규칙 (정규식, 대소문자 무시)", KW_PATTERN.pattern),
        ("", ""),
        ("입력하는 칸", "노란 칸만: '검증'(O = 판정 맞음, X = 판정 틀림, 드롭다운) · '메모'(자유)"),
        ("판정 틀림(X)의 뜻", "채택인데 K-water 아님 = 오탐 / 탈락인데 K-water 맞음 = 누락 / 소속 문자열은 판정 칸 기준"),
        ("추천 순서", "① 소속_문자열 (저자 수 많은 순) → ② 논문_탈락 → ③ 논문_채택 표본"),
        ("예시", "소속 'Korea Water Resources Association' 판정=K-water 라면 → 검증=X, 메모='학회(KWRA), 공사 아님'"),
    ]
    for a, b in lines:
        ws.append([a, b])
    for row in ws.iter_rows():
        for c in row:
            c.font = F
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A1"].font = Font(name="Arial", size=13, bold=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100

    # ── 소속 문자열 ──
    ws = wb.create_sheet("소속_문자열")
    table(ws, ["판정", "소속 문자열", "저자 수", "걸린 검색어", "검증(O/X)", "메모"], aff_rows,
          [10, 70, 9, 34, 10, 30], ["검증(O/X)", "메모"])

    # ── 논문 ──
    hdr = ["판정", "연도", "학술지", "국문 제목", "영문 제목", "K-water로 본 소속", "전체 저자 소속",
           "걸린 검색어", "URL", "검증(O/X)", "메모"]
    wid = [8, 7, 22, 48, 40, 40, 55, 26, 30, 10, 30]
    table(wb.create_sheet("논문_탈락"), hdr, [r for r in paper_rows if r[0] == "탈락"], wid, hdr[-2:])
    table(wb.create_sheet("논문_채택"), hdr, [r for r in paper_rows if r[0] == "채택"], wid, hdr[-2:])

    # ── 요약 (수식 — 검증 칸을 채우면 자동 집계) ──
    ws = wb.create_sheet("요약", 1)
    ws.append(["시트", "행 수", "검증 완료", "O (맞음)", "X (틀림)", "틀림 비율"])
    for i, (name, col) in enumerate([("소속_문자열", "E"), ("논문_탈락", "J"), ("논문_채택", "J")], start=2):
        q = f"'{name}'"
        ws.append([name, f"=COUNTA({q}!A:A)-1", f"=COUNTA({q}!{col}:{col})-1",
                   f'=COUNTIF({q}!{col}:{col},"O")', f'=COUNTIF({q}!{col}:{col},"X")',
                   f"=IF(C{i}=0,\"-\",E{i}/C{i})"])
        ws[f"F{i}"].number_format = "0.0%"
    for row in ws.iter_rows():
        for c in row:
            c.font = F
    for c in ws[1]:
        c.font, c.fill = FB, HEAD
    for col, w in zip("ABCDEF", [16, 10, 12, 12, 12, 12]):
        ws.column_dimensions[col].width = w
    ws["A6"] = "검증 완료 = 각 시트 '검증(O/X)' 칸을 채운 행 수 (머리글 제외). 수치는 노란 칸 입력에 따라 자동 갱신."
    ws["A6"].font = Font(name="Arial", size=9, italic=True)
    ws["A2"].comment = Comment("판정 기준과 입력 방법은 '안내' 시트 참고", "co-scientist-bot")
    wb.calculation.fullCalcOnLoad = True                         # 엑셀이 열 때 요약 수식을 전부 계산
    wb.save(path)


def main():
    if not KEY:
        print("❌ KCI_API_KEY 없음")
        return 1
    reg = load_registry()
    y0, y1 = int(reg.get("archive_from_year", 2000)), time.gmtime().tm_year
    papers, aff_n, aff_q = {}, Counter(), defaultdict(set)
    for y in range(y0, y1 + 1):
        for q in KW_QUERIES:
            kept, _, _ = fetch_year(None, None, y, {"affiliation": q})
            for r in kept:
                k = r["id"] or (r["title_ko"], r["journal"], r["volume"], r["issue"])
                p = papers.setdefault(k, {**r, "queries": set()})
                p["queries"].add(q)
                for a in r.get("authors", []):
                    if a.get("affil"):
                        aff_n[a["affil"]] += 1
                        aff_q[a["affil"]].add(q)
        print(f"   {y}: 누적 {len(papers)}편")

    OUT.mkdir(parents=True, exist_ok=True)
    with open(OUT / "kwater_papers.csv", "w", newline="", encoding="utf-8-sig") as f:   # 엑셀 한글 깨짐 방지
        w = csv.writer(f)
        w.writerow(["판정", "연도", "학술지", "국문 제목", "영문 제목", "K-water로 본 소속", "전체 저자 소속",
                    "걸린 검색어", "URL"])
        rows = []
        for p in papers.values():
            affs = [a.get("affil", "") for a in p.get("authors", [])]
            hit = [a for a in affs if KW_PATTERN.search(a)]
            rows.append(["채택" if hit else "탈락", p.get("year"), p.get("journal"), p.get("title_ko"),
                         p.get("title_en"), " | ".join(dict.fromkeys(hit)), " | ".join(dict.fromkeys(affs)),
                         ", ".join(sorted(p["queries"])), p.get("url")])
        rows.sort(key=lambda r: (r[0] != "탈락", str(r[1]), str(r[2])))                  # 탈락 먼저
        w.writerows(rows)
    with open(OUT / "kwater_affiliations.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["판정", "소속 문자열", "저자 수", "걸린 검색어"])
        w.writerows(sorted((["K-water" if KW_PATTERN.search(a) else "아님", a, n, ", ".join(sorted(aff_q[a]))]
                            for a, n in aff_n.items()), key=lambda r: (r[0], -r[2])))
    acc = sum(r[0] == "채택" for r in rows)
    aff_rows = sorted((["K-water" if KW_PATTERN.search(a) else "아님", a, n, ", ".join(sorted(aff_q[a]))]
                       for a, n in aff_n.items()), key=lambda r: (r[0], -r[2]))
    write_xlsx(OUT / "kwater_review.xlsx", rows, aff_rows,
               {"updated": time.strftime('%Y-%m-%d %H:%M UTC', time.gmtime())})
    (OUT / "README.md").write_text(f"""# K-water 소속 논문 분류 검증표

KCI(한국학술지인용색인) 데이터 활용 — **비공개 보관, 재배포 금지**. 갱신: {time.strftime('%Y-%m-%d %H:%M UTC', time.gmtime())}

- 검색어: {', '.join(KW_QUERIES)} (KCI `affiliation=` 검색)
- 재확인 규칙(정규식, 대소문자 무시): `{KW_PATTERN.pattern}`
- 검색에 걸린 논문 {len(rows):,}편 → **채택 {acc:,}편 · 탈락 {len(rows) - acc:,}편**

## 내 PC에서 검증
`kwater_review.xlsx`를 내려받아 엑셀로 여세요 (시트: 안내 · 요약 · 소속_문자열 · 논문_탈락 · 논문_채택).
노란 칸(검증 O/X 드롭다운, 메모)만 채우면 '요약' 시트가 자동 집계합니다.

## 보는 법
1. `kwater_affiliations.csv` — 소속 문자열별 빈도. '판정=K-water'인데 K-water가 아닌 것(오탐),
   '판정=아님'인데 K-water인 것(누락, 예: 표기 변형)을 찾으면 규칙을 고칩니다.
2. `kwater_papers.csv` — 논문 단위. **탈락이 먼저** 나옵니다. 탈락 논문의 '전체 저자 소속'에
   K-water 표기가 숨어 있으면 누락입니다. 채택 논문은 'K-water로 본 소속'이 맞는지 확인합니다.
""", encoding="utf-8")
    print(f"✅ 검증표: 검색 결과 {len(rows):,}편 → 채택 {acc:,} · 탈락 {len(rows) - acc:,} · 소속 문자열 {len(aff_n):,}종")
    return 0


if __name__ == "__main__":
    sys.exit(main())
